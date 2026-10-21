"""Step 1 of the Excel report: build excel/sales_report.xlsx with openpyxl.

Sheets: README, KPIs (formulas + conditional formatting + native charts), Segments (XLOOKUP), Data (order level), Items (item level).
Step 2 (excel/add_excel_features.py) adds PivotTables, slicers, Power Query and charts with Excel COM.

Run from the project root:  .venv\\Scripts\\python excel\\build_workbook.py
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.helpers import STATE_TO_REGION  # noqa: E402

PROC = ROOT / "data" / "processed"
OUT = ROOT / "sql" / "outputs"
XLSX = ROOT / "excel" / "sales_report.xlsx"

HEAD_FILL = PatternFill("solid", fgColor="1C5CAB")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=16)
NOTE_FONT = Font(italic=True, color="52514E")
FIRST_MONTH, LAST_MONTH = pd.Timestamp("2017-01-01").date(), pd.Timestamp("2018-08-01").date()   # full months
SAMPLE_SIZE = 500   # customers shown in the Segments list (random sample, seed 42)

SEGMENT_INFO = [
    ("Champions", "Bought recently, more than once, and spent a lot (recency score 4-5, 2+ orders, money score 4-5).", "Reward them. Ask for reviews and referrals."),
    ("Loyal", "Repeat buyers (2+ orders) with a recent order (recency score 3-5) who are not Champions.", "Keep them with offers and early access."),
    ("At Risk", "Repeat buyers (2+ orders) who have not bought for a long time (recency score 1-2).", "Send a win-back message now."),
    ("New Customers", "One order only, and it was very recent (recency score 4-5).", "Welcome them and push a second order."),
    ("Needs Attention", "One order only, medium recency (recency score 3).", "Remind them with a relevant offer."),
    ("About to Sleep", "One order only, older (recency score 2).", "Light reminder; do not spend much."),
    ("Lost", "One order only, the oldest group (recency score 1).", "Low priority. Test a cheap reactivation."),
]


def style_header(ws, row, first_col, last_col):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def add_table(ws, name, n_rows, n_cols):
    table = Table(displayName=name, ref=f"A1:{get_column_letter(n_cols)}{n_rows + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def write_df(ws, df, formats):
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    for col_name, fmt in formats.items():
        letter = get_column_letter(list(df.columns).index(col_name) + 1)
        for cell in ws[letter][1:]:
            cell.number_format = fmt
    ws.freeze_panes = "A2"


def load_data():
    orders = pd.read_parquet(PROC / "fact_orders.parquet")
    items = pd.read_parquet(PROC / "fact_order_items.parquet")
    products = pd.read_parquet(PROC / "dim_product.parquet")[["product_id", "category_en"]]
    valid = orders[orders.is_valid_sale].copy()
    valid["region"] = valid.customer_state.map(STATE_TO_REGION)
    data = pd.DataFrame({
        "order_id": valid.order_id,
        "purchase_date": valid.purchase_ts.dt.date,
        "month": valid.purchase_ts.dt.to_period("M").dt.to_timestamp().dt.date,
        "customer_state": valid.customer_state,
        "region": valid.region,
        "payment_type": valid.main_payment_type.fillna("unknown"),
        "items_value": valid.items_value.astype(float),
        "freight_value": valid.freight_value.astype(float),
        "review_score": valid.review_score.astype("float64"),
        "delivery_days": valid.delivery_days.astype("float64"),
        "is_late": valid.is_late.astype("float64"),
    }).reset_index(drop=True)
    data["full_month"] = ["Yes" if FIRST_MONTH <= m <= LAST_MONTH else "No" for m in data["month"]]
    data = data.astype(object).where(data.notna(), None)

    it = items.merge(valid[["order_id", "customer_state", "region"]], on="order_id", how="inner").merge(products, on="product_id", how="left")
    month = pd.to_datetime(it.purchase_date_key.astype(str), format="%Y%m%d").dt.to_period("M").dt.to_timestamp()
    items_x = pd.DataFrame({
        "month": month.dt.date, "year": month.dt.year, "region": it.region, "state": it.customer_state,
        "category": it.category_en, "price": it.price.astype(float), "freight_value": it.freight_value.astype(float),
    })
    items_x["full_month"] = ["Yes" if FIRST_MONTH <= m <= LAST_MONTH else "No" for m in items_x["month"]]
    return data, items_x, valid


def build_kpis(ws, data_rows, months):
    ws["A1"] = "Sales KPIs (valid sales: not canceled, not unavailable)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Revenue = item price without freight. All numbers are formulas on the Data sheet."
    ws["A2"].font = NOTE_FONT

    headline = [
        ("Total revenue (all months)", "=SUM(tblData[items_value])", "#,##0.00"),
        ("Total orders (all months)", "=ROWS(tblData[order_id])", "#,##0"),
        ("Average order value (AOV)", "=IFERROR(B4/B5,0)", "#,##0.00"),
        ("Average review score", "=IFERROR(AVERAGE(tblData[review_score]),0)", "0.00"),
        ("Late delivery rate (delivered orders)", "=IFERROR(AVERAGE(tblData[is_late]),0)", "0.0%"),
    ]
    for i, (label, formula, fmt) in enumerate(headline):
        r = 4 + i
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        c = ws.cell(row=r, column=2, value=formula)
        c.number_format = fmt

    header_row = 11
    heads = ["Month", "Orders", "Revenue", "AOV", "MoM growth", "Freight", "Avg review score", "Late delivery rate"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=header_row, column=j, value=h)
    style_header(ws, header_row, 1, len(heads))
    ws.cell(row=header_row - 1, column=1, value="Monthly KPIs, full months 2017-01 to 2018-08").font = Font(bold=True)

    for i, m in enumerate(months):
        r = header_row + 1 + i
        ws.cell(row=r, column=1, value=m).number_format = "yyyy-mm"
        ws.cell(row=r, column=2, value=f"=COUNTIFS(tblData[month],$A{r})").number_format = "#,##0"
        ws.cell(row=r, column=3, value=f"=SUMIFS(tblData[items_value],tblData[month],$A{r})").number_format = "#,##0.00"
        ws.cell(row=r, column=4, value=f"=IFERROR(C{r}/B{r},0)").number_format = "#,##0.00"
        ws.cell(row=r, column=5, value=("" if i == 0 else f'=IFERROR(C{r}/C{r - 1}-1,"")')).number_format = "0.0%"
        ws.cell(row=r, column=6, value=f"=SUMIFS(tblData[freight_value],tblData[month],$A{r})").number_format = "#,##0.00"
        ws.cell(row=r, column=7, value=f'=IFERROR(AVERAGEIFS(tblData[review_score],tblData[month],$A{r}),"")').number_format = "0.00"
        ws.cell(row=r, column=8, value=f'=IFERROR(AVERAGEIFS(tblData[is_late],tblData[month],$A{r}),"")').number_format = "0.0%"
    first, last = header_row + 1, header_row + len(months)

    # Conditional formatting: growth up (green) / down (red); review score below the overall average
    green = PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")
    red = PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE")
    ws.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(operator="greaterThan", formula=["0"], fill=green, font=Font(color="006100")))
    ws.conditional_formatting.add(f"E{first}:E{last}", CellIsRule(operator="lessThan", formula=["0"], fill=red, font=Font(color="9C0006")))
    ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(formula=[f'AND(ISNUMBER(G{first}),G{first}<$B$7-0.1)'], fill=red, font=Font(color="9C0006")))
    ws.cell(row=last + 1, column=1, value="Green/red in MoM growth = up/down. Red review score = more than 0.1 below the overall average (B7).").font = NOTE_FONT

    for col, width in zip("ABCDEFGH", [38, 18, 16, 12, 13, 14, 16, 18]):
        ws.column_dimensions[col].width = width

    # The 3 native charts of this sheet are added by excel/add_excel_features.py (Excel COM), which gives cleaner charts.


def build_segments(ws):
    rfm = pd.read_csv(next(OUT.glob("q12_*.csv")))
    order = [s[0] for s in SEGMENT_INFO]
    summary = rfm.groupby("segment").agg(customers=("customer_unique_id", "count"), avg_recency=("recency_days", "mean"),
                                         avg_orders=("frequency", "mean"), avg_money=("monetary", "mean"),
                                         total_money=("monetary", "sum")).reindex(order).reset_index()

    ws["A1"] = "RFM customer segments (from query Q12)"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Segment table: all customers (values from Q12). The description columns use XLOOKUP on the segment info table (columns K to M)."
    ws["A2"].font = NOTE_FONT
    heads = ["Segment", "Customers", "Share of customers", "Avg days since last order", "Avg orders", "Avg revenue per customer",
             "Total revenue", "Description (XLOOKUP)"]
    for j, h in enumerate(heads, start=1):
        ws.cell(row=4, column=j, value=h)
    style_header(ws, 4, 1, len(heads))
    n = len(summary)
    for i, r in summary.iterrows():
        row = 5 + i
        ws.cell(row=row, column=1, value=r.segment)
        ws.cell(row=row, column=2, value=int(r.customers)).number_format = "#,##0"
        ws.cell(row=row, column=3, value=f"=B{row}/$B${5 + n}").number_format = "0.0%"
        ws.cell(row=row, column=4, value=round(float(r.avg_recency), 1)).number_format = "0.0"
        ws.cell(row=row, column=5, value=round(float(r.avg_orders), 3)).number_format = "0.000"
        ws.cell(row=row, column=6, value=round(float(r.avg_money), 2)).number_format = "#,##0.00"
        ws.cell(row=row, column=7, value=round(float(r.total_money), 2)).number_format = "#,##0.00"
        ws.cell(row=row, column=8, value=f'=_xlfn.XLOOKUP($A{row},tblSegmentInfo[Segment],tblSegmentInfo[Description],"not found")')
    tot = 5 + n
    ws.cell(row=tot, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=tot, column=2, value=f"=SUM(B5:B{tot - 1})").number_format = "#,##0"
    ws.cell(row=tot, column=3, value=f"=SUM(C5:C{tot - 1})").number_format = "0.0%"
    ws.cell(row=tot, column=7, value=f"=SUM(G5:G{tot - 1})").number_format = "#,##0.00"

    # Segment info table (lookup source)
    for j, h in enumerate(["Segment", "Description", "Suggested action"], start=11):
        ws.cell(row=4, column=j, value=h)
    for i, (seg, desc, act) in enumerate(SEGMENT_INFO):
        ws.cell(row=5 + i, column=11, value=seg)
        ws.cell(row=5 + i, column=12, value=desc)
        ws.cell(row=5 + i, column=13, value=act)
    info = Table(displayName="tblSegmentInfo", ref=f"K4:M{4 + len(SEGMENT_INFO)}")
    info.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(info)

    # Sample of customers with XLOOKUP
    start = tot + 3
    ws.cell(row=start - 1, column=1, value=f"Sample of {SAMPLE_SIZE} customers (random, seed 42; the full list is in sql/outputs/q12_rfm_segments.csv)").font = Font(bold=True)
    cols = ["Customer id", "State", "Recency (days)", "Orders", "Revenue", "RFM score", "Segment", "Description (XLOOKUP)", "", "", "Suggested action (XLOOKUP)"]
    for j, h in enumerate(cols, start=1):
        if h:
            ws.cell(row=start, column=j, value=h)
    style_header(ws, start, 1, 8)
    ws.cell(row=start, column=11).fill, ws.cell(row=start, column=11).font = HEAD_FILL, HEAD_FONT
    sample = rfm.sample(SAMPLE_SIZE, random_state=42).sort_values("monetary", ascending=False)
    for i, r in enumerate(sample.itertuples(index=False)):
        row = start + 1 + i
        ws.cell(row=row, column=1, value=r.customer_unique_id)
        ws.cell(row=row, column=2, value=r.customer_state)
        ws.cell(row=row, column=3, value=int(r.recency_days))
        ws.cell(row=row, column=4, value=int(r.frequency))
        ws.cell(row=row, column=5, value=float(r.monetary)).number_format = "#,##0.00"
        ws.cell(row=row, column=6, value=str(r.rfm_score))
        ws.cell(row=row, column=7, value=r.segment)
        ws.cell(row=row, column=8, value=f'=_xlfn.XLOOKUP($G{row},tblSegmentInfo[Segment],tblSegmentInfo[Description],"not found")')
        ws.cell(row=row, column=11, value=f'=_xlfn.XLOOKUP($G{row},tblSegmentInfo[Segment],tblSegmentInfo[Suggested action],"not found")')
    for col, width in zip("ABCDEFGHIJKLM", [36, 11, 14, 11, 14, 11, 16, 60, 4, 4, 38, 70, 42]):
        ws.column_dimensions[col].width = width


def build_readme(wb, ws, n_data, n_items):
    lines = [
        ("Ecommerce_Sales_Analytics: Excel report", TITLE_FONT),
        ("Part of a data analyst portfolio project. Data: Olist Brazilian E-Commerce (2016-2018), cleaned in this project.", None),
        ("", None),
        ("Sheets", Font(bold=True)),
        ("KPIs: revenue, orders, AOV, MoM growth, freight, review score and late rate per month. Formulas on the Data sheet, conditional formatting, 3 native charts (added in step 2).", None),
        ("Segments: RFM segments from query Q12, with XLOOKUP to a description table.", None),
        ("Pivot: PivotTables with slicers (region, year) and PivotCharts (built with Excel automation). Revenue by month shows full months only.", None),
        ("PQ_Top3_Products, PQ_State_Orders: Power Query queries that import the SQL result files Q06 and Q15.", None),
        (f"Data: one row per valid order ({n_data:,} rows).", None),
        (f"Items: one row per item sold in a valid order ({n_items:,} rows).", None),
        ("", None),
        ("Definitions", Font(bold=True)),
        ("Valid sale: the order has items and its status is not canceled or unavailable.", None),
        ("Revenue: sum of item prices, without freight. AOV = revenue / orders.", None),
        ("MoM growth: revenue of the month divided by revenue of the month before, minus 1.", None),
        ("Full months: 2017-01 to 2018-08. Other months have too few orders (see notebook 01, decision 11).", None),
        ("", None),
        ("Notes", Font(bold=True)),
        (f"The Segments sheet shows a random sample of {SAMPLE_SIZE} customers; the segment summary uses all customers.", None),
        ("The Power Query queries read the SQL result CSV files from the folder in the named cell DataFolder (yellow cell at the bottom of this sheet).", None),
        ("If you move the project, type the new folder in that cell and use Data > Refresh All. No query has to be edited.", None),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if font:
            c.font = font
    # Parameter for Power Query: a named cell with the folder of the SQL result files
    row = len(lines) + 2
    ws.cell(row=row, column=1, value="Power Query parameter: DataFolder (folder with the SQL result CSV files)").font = Font(bold=True)
    cell = ws.cell(row=row + 1, column=1, value=str(OUT))
    cell.fill = PatternFill("solid", fgColor="FFF2CC")
    cell.font = Font(bold=True)
    wb.defined_names["DataFolder"] = DefinedName("DataFolder", attr_text=f"README!$A${row + 1}")
    ws.column_dimensions["A"].width = 150


def main():
    data, items, valid = load_data()
    months = sorted(m for m in set(data["month"]) if FIRST_MONTH <= m <= LAST_MONTH)

    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "README"
    ws_kpi = wb.create_sheet("KPIs")
    ws_seg = wb.create_sheet("Segments")
    ws_data = wb.create_sheet("Data")
    ws_items = wb.create_sheet("Items")

    write_df(ws_data, data, {"purchase_date": "yyyy-mm-dd", "month": "yyyy-mm", "items_value": "#,##0.00",
                             "freight_value": "#,##0.00", "review_score": "0", "delivery_days": "0.0", "is_late": "0"})
    add_table(ws_data, "tblData", len(data), data.shape[1])
    for col, width in zip("ABCDEFGHIJKL", [34, 14, 10, 14, 14, 14, 12, 13, 13, 13, 8, 11]):
        ws_data.column_dimensions[col].width = width
    # Data sheet: low review scores (1 or 2) in light red
    ws_data.conditional_formatting.add(f"I2:I{len(data) + 1}", CellIsRule(operator="between", formula=["1", "2"],
                                       fill=PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE"), font=Font(color="9C0006")))

    write_df(ws_items, items, {"month": "yyyy-mm", "price": "#,##0.00", "freight_value": "#,##0.00"})
    add_table(ws_items, "tblItems", len(items), items.shape[1])
    for col, width in zip("ABCDEFGH", [10, 8, 14, 8, 28, 12, 13, 11]):
        ws_items.column_dimensions[col].width = width

    build_kpis(ws_kpi, len(data), months)
    build_segments(ws_seg)
    build_readme(wb, ws_readme, len(data), len(items))

    XLSX.parent.mkdir(exist_ok=True)
    wb.save(XLSX)
    print(f"saved {XLSX.relative_to(ROOT)}  ({XLSX.stat().st_size / 1e6:.1f} MB); Data rows: {len(data):,}; Items rows: {len(items):,}")


if __name__ == "__main__":
    main()
